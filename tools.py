import datetime
import os
from openpyxl import Workbook, load_workbook

def get_current_time():
    """Возвращает текущую дату и время"""

    now = datetime.datetime.now()
    return now.strftime("Сегодня: %Y-%m-%d, %A. Время: %H:%M")


def save_tasks_to_matrix(tasks):
    """
    Сохраняет список задач в Excel-файл matrix.xlsx.
    Если файл не существует, он будет создан с заголовками.
    """
    filename = "matrix.xlsx"
    headers = ["Дата добавления", "Задача", "Квадрант", "Обоснование"]

    # 1. Проверяем, существует ли файл, и открываем/создаем его
    if os.path.exists(filename):
        wb = load_workbook(filename)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Задачи"
        ws.append(headers)  # Добавляем шапку, если файл новый

    # 2. Добавляем данные
    current_timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

    print("\n--- ЗАПИСЬ В EXCEL ---")
    for task in tasks:
        title = task.get("title", "Без названия")
        quadrant = task.get("quadrant", "?")
        reason = task.get("reason", "Причина не указана")

        # Добавляем строку в Excel
        ws.append([current_timestamp, title, quadrant, reason])
        print(f"Записано: {title} (Квадрант {quadrant})")

    # 3. Сохраняем файл
    try:
        wb.save(filename)
        return f"Успешно: {len(tasks)} задач сохранено в файл {filename}."
    except Exception as e:
        return f"Ошибка при сохранении файла: {str(e)}"


TOOLS_DESCRIPTION = [
    {
        "name": "get_current_time",
        "description": "Позволяет узнать текущую дату, день недели и время. Используй это, чтобы правильно рассчитать дедлайны задач (например, если пользователь говорит 'завтра' или 'в среду').",
        "parameters": {
            "type": "object",
            "properties": {}
        }
    },
    {
        "name": "save_tasks_to_matrix",
        "description": "Сохраняет классифицированные задачи в матрицу Эйзенхауэра.",
        "parameters": {
            "type": "object",
            "properties": {
                "tasks": {
                    "type": "array",
                    "description": "Список задач для сохранения",
                    "items": {
                        "type": "object",
                        "properties": {
                            "title": {
                                "type": "string",
                                "description": "Краткое название задачи (например, 'Купить корм')"
                            },
                            "quadrant": {
                                "type": "integer",
                                "description": "Номер квадранта (1-Срочно/Важно, 2-Важно/Не срочно, 3-Срочно/Не важно, 4-Не важно/Не срочно)",
                                "enum": [1, 2, 3, 4]
                            },
                            "reason": {
                                "type": "string",
                                "description": "Логика, почему задача попала в этот квадрант (например, 'Дедлайн через 2 дня, работа')"
                            }
                        },
                        "required": ["title", "quadrant", "reason"]
                    }
                }
            },
            "required": ["tasks"]
        }
    }
]